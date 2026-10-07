import io
import unittest
import zipfile
import xml.etree.ElementTree as ET

import numpy as np
import openpyxl
import pandas as pd

from inventory_structure_analysis import (
    DETAIL_COLUMNS,
    SHEET_NAMES,
    SUMMARY_AMOUNT_COLUMN,
    SUMMARY_SHARE_COLUMN,
    InventoryStructureAnalyzer,
    composition_chart_color,
)


REPLENISHMENT_COLUMNS = pd.MultiIndex.from_tuples(
    [
        ("物料", "物料"),
        ("物料", "物料描述"),
        ("产品基本信息", "产品状态"),
        ("产品基本信息", "产品组"),
        ("产品基本信息", "产品组描述"),
        ("产品基本信息", "国贸产品分类码"),
        ("产品基本信息", "级别"),
        ("产品基本信息", "Current Stock level"),
        ("产品基本信息", "库存月数"),
        ("库存信息", "当前可用库存"),
        ("在途在产信息", "在途+在产"),
        ("预测销量", "月预测"),
        ("预测销量", "月预测（修改）"),
        ("服务水平信息", "全供应链存销比"),
    ]
)


def replenishment(rows):
    return pd.DataFrame(rows, columns=REPLENISHMENT_COLUMNS)


def row(code, stock=100, forecast=10, revised=np.nan, target=6, status="Active", group="G1", group_desc="Group 1", category="C1", level="Parent"):
    return [code, f"Item {code}", status, group, group_desc, category, level, "RA", target, stock, 5, forecast, revised, 12]


class InventoryStructureAnalyzerTests(unittest.TestCase):
    def setUp(self):
        self.analyzer = InventoryStructureAnalyzer()

    def analyze(self, rows, prices=None):
        prices = prices if prices is not None else pd.DataFrame({"物料代码": [r[0] for r in rows], "单价": [2] * len(rows)})
        return self.analyzer.analyze(replenishment(rows), prices)

    def test_revised_forecast_has_priority(self):
        result = self.analyze([row("A", stock=90, forecast=10, revised=30)])
        self.assertEqual(result.detail.loc[0, "有效预测月销"], 30)
        self.assertEqual(result.detail.loc[0, "可用库存存销比"], 3)
        self.assertAlmostEqual(result.detail.loc[0, "全供应链存销比"], 95 / 30)

    def test_price_columns_are_auto_detected_only_from_known_names(self):
        recognized = pd.DataFrame({"物料代码": ["A"], "销售组织": [8374], "单价": [2.5]})
        self.assertEqual(self.analyzer.detect_price_columns(recognized), ("物料代码", "单价"))

        unrecognized = pd.DataFrame({"物料代码": ["A"], "销售组织": [8374]})
        self.assertEqual(self.analyzer.detect_price_columns(unrecognized), ("物料代码", None))

    def test_new_replenishment_headers_use_l1_to_l4_without_full_chain_ratio(self):
        source = replenishment([row("A", stock=90, forecast=15), row("B", stock=120, forecast=20)])
        source = source.drop(
            columns=[
                ("产品基本信息", "国贸产品分类码"),
                ("服务水平信息", "全供应链存销比"),
            ]
        )
        for level, values in {
            "L1": ["一级A", "一级B"],
            "L2": ["二级A", "二级B"],
            "L3": ["三级A", "三级B"],
            "L4": ["四级A", "四级B"],
        }.items():
            source[("实际销量", f"国贸产品分类码{level}")] = values

        prices = pd.DataFrame({"物料代码": ["A", "B"], "单价": [2, 2]})
        result = self.analyzer.analyze(source, prices)

        self.assertEqual(
            result.diagnostics["国贸分类维度"],
            ("国贸产品分类码L1", "国贸产品分类码L2", "国贸产品分类码L3", "国贸产品分类码L4"),
        )
        self.assertEqual(result.detail["全供应链存销比"].tolist(), [95 / 15, 125 / 20])
        self.assertEqual(result.detail["可用库存存销比"].tolist(), [6.0, 6.0])
        for level in ["L2", "L3", "L4"]:
            self.assertIn(f"国贸分类码{level}前80%", result.summaries)
            self.assertIn(f"国贸分类码{level}明细", result.summaries)
            self.assertEqual(
                result.diagnostics["国贸分类诊断"][f"国贸产品分类码{level}"]["国贸分类码分组总数"],
                2,
            )
        self.assertIn("国贸分类码L1汇总", result.summaries)
        self.assertNotIn("国贸分类码L1分层", result.summaries)
        self.assertNotIn("国贸分类码L1明细", result.summaries)

        workbook = openpyxl.load_workbook(io.BytesIO(self.analyzer.export_excel(result)))
        self.assertEqual(workbook.sheetnames, SHEET_NAMES)
        for level in range(1, 5):
            self.assertIn(f"国贸产品分类码L{level}", [
                cell.value for cell in workbook[f"新品推荐L{level}"][1]
            ])

    def test_l2_top_eighty_overview_uses_categories_not_amount_tiers(self):
        source = replenishment(
            [
                row("A", stock=40_000_000),
                row("B", stock=30_000_000),
                row("C", stock=20_000_000),
                row("D", stock=10_000_000),
            ]
        ).drop(columns=[("产品基本信息", "国贸产品分类码")])
        source[("实际销量", "国贸产品分类码L2")] = ["二级A", "二级B", "二级C", "二级D"]
        prices = pd.DataFrame({"物料代码": list("ABCD"), "单价": [1, 1, 1, 1]})

        result = self.analyzer.analyze(source, prices)
        overview = result.summaries["国贸分类码L2前80%"]
        detail = result.summaries["国贸分类码L2明细"]

        self.assertEqual(overview["国贸产品分类码L2"].tolist(), ["二级A", "二级B", "二级C"])
        self.assertAlmostEqual(overview.iloc[-1]["累计占比"], 0.9)
        self.assertNotIn("库存分层分类", overview.columns)
        self.assertNotIn("库存分层分类", detail.columns)
        self.assertEqual(
            result.diagnostics["国贸分类诊断"]["国贸产品分类码L2"]["国贸分类码分组总数"],
            4,
        )

    def test_only_parent_level_is_included_in_all_analysis(self):
        result = self.analyze([row("A", level=" parent "), row("B", level="SON"), row("C", level=np.nan)])
        self.assertEqual(result.detail["物料"].tolist(), ["A"])
        self.assertEqual(result.diagnostics["源文件记录数"], 3)
        self.assertEqual(result.diagnostics["纳入Parent记录数"], 1)
        self.assertEqual(result.diagnostics["忽略非Parent记录数"], 2)
        self.assertEqual(result.summaries["按产品状态汇总"].iloc[-1]["型号数量"], 1)

    def test_sales_organization_is_first_column_in_inventory_excel_sheets(self):
        source = replenishment([row("A", stock=250)])
        source[("快照日期", "销售组织")] = ["EU01"]
        result = self.analyzer.analyze(
            source, pd.DataFrame({"物料代码": ["A"], "单价": [2]})
        )

        workbook = openpyxl.load_workbook(io.BytesIO(self.analyzer.export_excel(result)))
        for sheet_name in SHEET_NAMES[:3]:
            self.assertEqual(workbook[sheet_name]["A1"].value, "销售组织")
        self.assertEqual(workbook["库存24个月以上"]["A2"].value, "EU01")

    def test_original_forecast_used_when_revised_is_empty(self):
        result = self.analyze([row("A", stock=90, forecast=15)])
        self.assertEqual(result.detail.loc[0, "有效预测月销"], 15)
        self.assertEqual(result.detail.loc[0, "可用库存存销比"], 6)

    def test_exactly_24_months_is_long_term(self):
        result = self.analyze([row("A", stock=240, forecast=10)])
        self.assertEqual(result.detail.loc[0, "库存区间"], "库存24个月以上")

    def test_exactly_target_plus_three_is_normal(self):
        result = self.analyze([row("A", stock=90, forecast=10, target=6)])
        self.assertEqual(result.detail.loc[0, "库存区间"], "正常及关注库存")

    def test_intervals_are_mutually_exclusive_and_complete_for_valid_rows(self):
        result = self.analyze([row("A", 250), row("B", 100), row("C", 90)])
        codes = []
        for name in ["库存24个月以上", "超目标库存", "正常及关注库存"]:
            codes.extend(result.views[name]["物料"].tolist())
        self.assertCountEqual(codes, ["A", "B", "C"])
        self.assertEqual(len(codes), len(set(codes)))

    def test_nonpositive_forecast_does_not_divide_by_zero(self):
        result = self.analyze([row("A", forecast=0), row("B", forecast=-2)])
        self.assertTrue(result.detail["可用库存存销比"].isna().all())
        self.assertTrue(result.detail["全供应链存销比"].isna().all())
        self.assertTrue(result.detail["计算状态"].str.contains("无有效销量").all())
        self.assertTrue(result.detail["库存区间"].eq("数据问题").all())

    def test_missing_forecast_and_stock_are_not_coerced_to_zero(self):
        result = self.analyze([row("A", forecast=np.nan), row("B", stock=np.nan)])
        self.assertIn("预测月销缺失", result.detail.loc[0, "计算状态"])
        self.assertIn("可用库存缺失", result.detail.loc[1, "计算状态"])
        self.assertTrue(result.detail["可用库存存销比"].isna().all())

    def test_missing_price(self):
        result = self.analyze([row("A")], pd.DataFrame({"物料代码": ["B"], "单价": [2]}))
        self.assertTrue(pd.isna(result.detail.loc[0, "单价"]))
        self.assertTrue(pd.isna(result.detail.loc[0, "可用库存金额"]))
        self.assertIn("价格缺失", result.detail.loc[0, "计算状态"])
        self.assertEqual(result.diagnostics["价格缺失数量"], 1)

    def test_no_sales_and_missing_stock_level_views_ignore_price_issues(self):
        rows = [row("A", forecast=0), row("B"), row("C")]
        rows[1][7] = np.nan
        result = self.analyze(rows, pd.DataFrame({"物料代码": ["A", "B"], "单价": [2, 2]}))
        self.assertEqual(result.views["无销量"]["物料"].tolist(), ["A"])
        self.assertEqual(result.views["无库存等级"]["物料"].tolist(), ["B"])
        self.assertIn("价格缺失", result.detail.loc[2, "计算状态"])
        self.assertNotIn("C", result.views["无销量"]["物料"].tolist())
        self.assertNotIn("C", result.views["无库存等级"]["物料"].tolist())

    def test_status_analyses_use_trade_category_and_abc_inventory_bands(self):
        source = replenishment([
            row("A", stock=240, status="NORMAL", group_desc="Group 1", category="Trade 1"),
            row("B", stock=100, status="PHASING OUT", group_desc="Group 2", category="Trade 2"),
            row("C", stock=90, status="NEW", group_desc="Group 1", category="Trade 1"),
        ])
        source[("产品基本信息", "泛欧ABC")] = ["A", "B", "C"]
        prices = pd.DataFrame({"物料代码": ["A", "B", "C"], "单价": [2, 2, 2]})
        result = self.analyzer.analyze(source, prices)

        trade = self.analyzer.summarize_status_trade(result.detail, "NEW", "国贸产品分类码")
        self.assertEqual(trade["国贸产品分类码"].tolist(), ["Trade 1"])
        self.assertEqual(trade["型号数量"].tolist(), [1])
        self.assertEqual(self.analyzer.available_product_groups(result.detail, "NORMAL"), ["Group 1"])
        inventory = self.analyzer.summarize_status_inventory(result.detail, "NORMAL", ["Group 1"])
        self.assertEqual(inventory.loc[0, "产品组描述"], "Group 1")
        self.assertEqual(inventory.loc[0, "ABC等级"], "A")
        self.assertEqual(inventory.loc[0, "库存区间"], "库存24个月以上")
        self.assertEqual(inventory.loc[0, "型号数量"], 1)

    def test_status_groups_filter_product_groups_and_keep_all_trade_levels(self):
        source = replenishment([
            row("A", status="NEW RELEASE (RECOMMEND)", group_desc="Group 1"),
            row("B", status="NEW", group_desc="Group 2"),
            row("C", status="NORMAL", group_desc="Group 1"),
            row("D", status="PHASING OUT", group_desc="Group 2"),
            row("E", status="EOL", group_desc="Group 1"),
        ])
        source[("产品基本信息", "泛欧ABC")] = ["A", "A", "B", "C", "C"]
        for level in range(1, 5):
            source[("产品基本信息", f"国贸产品分类码L{level}")] = [
                f"L{level}-{code}" for code in "ABCDE"
            ]
        prices = pd.DataFrame({"物料代码": list("ABCDE"), "单价": [2] * 5})
        result = self.analyzer.analyze(source, prices)

        self.assertEqual(result.diagnostics["国贸分类维度"], tuple(
            f"国贸产品分类码L{level}" for level in range(1, 5)
        ))
        self.assertEqual(self.analyzer.available_product_groups(result.detail, "NEW RELEASE (RECOMMEND)"), ["Group 1", "Group 2"])
        status_overview = self.analyzer.summarize_status_overview(result.detail)
        self.assertEqual(status_overview["产品状态类别"].tolist(), [
            "NEW RELEASE (RECOMMEND)", "NORMAL/PHASING OUT", "EOL", "合计"
        ])
        self.assertAlmostEqual(status_overview.iloc[:-1][SUMMARY_SHARE_COLUMN].sum(), 1)
        for level in range(1, 5):
            dimension = f"国贸产品分类码L{level}"
            summary = self.analyzer.summarize_status_trade(
                result.detail, "NEW RELEASE (RECOMMEND)", dimension, ["Group 1"]
            )
            self.assertEqual(summary[dimension].tolist(), [f"L{level}-A"])
        combined = self.analyzer.summarize_status_inventory(
            result.detail, "NORMAL/PHASING OUT", ["Group 1", "Group 2"]
        )
        self.assertEqual(set(combined["ABC等级"]), {"B", "C"})
        eol = self.analyzer.summarize_status_trade(result.detail, "EOL", "国贸产品分类码L4", ["Group 1"])
        self.assertEqual(eol["国贸产品分类码L4"].tolist(), ["L4-E"])
        workbook = openpyxl.load_workbook(io.BytesIO(self.analyzer.export_excel(result)))
        self.assertEqual(workbook.sheetnames, SHEET_NAMES)
        abc_sheet = workbook["正常渐退ABC"]
        self.assertEqual(len(abc_sheet._charts), 2)
        self.assertEqual([abc_sheet["A1"].value, abc_sheet["B1"].value], ["产品组", "产品组描述"])
        self.assertEqual(abc_sheet["F2"].number_format, "#,##0.00")
        self.assertEqual(abc_sheet._charts[1].anchor._from.row - abc_sheet._charts[0].anchor._from.row, 19)
        self.assertEqual(len(workbook["新品推荐L4"]._charts), 1)
        self.assertEqual(len(workbook["停产L4"]._charts), 1)
        self.assertNotIn("产品组分类透视图", workbook.sheetnames)
        self.assertEqual(workbook["新品推荐L4"]._charts[0].grouping, "percentStacked")

    def test_duplicate_material_with_same_price_is_deduplicated(self):
        prices = pd.DataFrame({"物料代码": ["A", "A"], "单价": [2, 2]})
        result = self.analyze([row("A")], prices)
        self.assertEqual(result.detail.loc[0, "单价"], 2)
        self.assertNotIn("价格冲突", result.detail.loc[0, "计算状态"])

    def test_duplicate_material_with_conflicting_prices_is_flagged(self):
        prices = pd.DataFrame({"物料代码": ["A", "A"], "单价": [2, 3]})
        result = self.analyze([row("A")], prices)
        self.assertTrue(pd.isna(result.detail.loc[0, "单价"]))
        self.assertIn("价格冲突", result.detail.loc[0, "计算状态"])
        self.assertEqual(result.diagnostics["价格冲突数量"], 1)

    def test_model_counts_are_distinct_material_codes(self):
        result = self.analyze([row("A"), row("A"), row("B")])
        summary = result.summaries["按产品状态汇总"]
        self.assertEqual(summary.iloc[-1]["型号数量"], 2)

    def test_three_summary_totals_match_detail_amount(self):
        result = self.analyze([row("A", stock=10), row("B", stock=20), row("C", stock=30)])
        detail_total = result.detail["可用库存金额"].sum(min_count=1)
        for name in ["按产品状态汇总", "按产品组汇总"]:
            summary = result.summaries[name]
            self.assertEqual(summary.iloc[-1][SUMMARY_AMOUNT_COLUMN], detail_total / 1_000_000)
            self.assertAlmostEqual(summary.iloc[:-1][SUMMARY_SHARE_COLUMN].sum(), 1.0)
            self.assertEqual(summary.iloc[-1][SUMMARY_SHARE_COLUMN], 1.0)
        trade_detail = result.summaries["按国贸分类码汇总"]
        self.assertEqual(trade_detail[SUMMARY_AMOUNT_COLUMN].sum(), detail_total / 1_000_000)
        self.assertAlmostEqual(trade_detail[SUMMARY_SHARE_COLUMN].sum(), 1.0)
        self.assertTrue(result.diagnostics["汇总金额核对一致"])

    def test_inventory_band_amounts_reconcile_with_total_without_missing_prices(self):
        rows = [
            row("A", stock=240),
            row("B", stock=100),
            row("C", stock=90),
            row("D", stock=100, forecast=0),
            row("E", stock=90),
        ]
        prices = pd.DataFrame({"物料代码": ["A", "B", "C", "D"], "单价": [2, 2, 2, 2]})
        result = self.analyze(rows, prices)
        amounts = result.diagnostics["库存区间金额"]
        self.assertEqual(amounts, {
            "库存24个月以上": 480,
            "超目标库存": 200,
            "正常及关注库存": 180,
            "数据问题": 200,
        })
        self.assertEqual(sum(amounts.values()), result.diagnostics["有效明细金额"])
        self.assertAlmostEqual(sum(value / result.diagnostics["有效明细金额"] for value in amounts.values()), 1)
        self.assertEqual(result.diagnostics["正常及关注库存数量"], 2)
        overview = self.analyzer.summarize_inventory_bands(result)
        self.assertAlmostEqual(overview.iloc[:-1][SUMMARY_AMOUNT_COLUMN].sum(), overview.iloc[-1][SUMMARY_AMOUNT_COLUMN])
        self.assertAlmostEqual(overview.iloc[:-1][SUMMARY_SHARE_COLUMN].sum(), 1)

    def test_trade_category_tables_include_tiers_and_single_detail(self):
        rows = [
            row("A", stock=2_000_000, category="C1"),
            row("B", stock=750_000, category="C2"),
            row("C", stock=200_000, category="C3"),
            row("D", stock=1_000, category="C4"),
            row("E", stock=-10_000, category="C5"),
            row("F", stock=np.nan, category="C6"),
        ]
        prices = pd.DataFrame({"物料代码": list("ABCDEF"), "单价": [1] * 6})
        result = self.analyze(rows, prices)
        overview = result.summaries["国贸分类码分层"]
        detail = result.summaries["按国贸分类码汇总"]

        self.assertNotIn("国贸重点分类码", result.summaries)
        self.assertNotIn("国贸数据问题", result.summaries)
        self.assertIn("单型号库存金额（k USD）", detail.columns)
        self.assertIn("累计占比", detail.columns)
        self.assertIn("库存分层分类", detail.columns)
        self.assertIn("负库存金额", detail["数据状态"].tolist())
        self.assertIn("金额缺失", detail["数据状态"].tolist())
        self.assertIn("金额接近零", detail["数据状态"].tolist())
        self.assertEqual(
            overview.iloc[:-1]["库存分层分类"].tolist(),
            [
                "核心（≥1 M USD）",
                "重点（0.5–1 M USD）",
                "一般（0.1–0.5 M USD）",
                "长尾（0–0.1 M USD）",
                "负金额",
            ],
        )
        self.assertEqual(overview.iloc[-1]["库存分层分类"], "合计")

    def test_missing_categories_are_labeled_unclassified(self):
        result = self.analyze([row("A", status=np.nan, group=np.nan, group_desc=np.nan, category=np.nan)])
        self.assertEqual(result.summaries["按产品状态汇总"].iloc[0]["产品状态"], "未分类")
        self.assertEqual(result.summaries["按国贸分类码汇总"].iloc[0]["国贸产品分类码"], "未分类")

    def test_excel_has_required_sheets_column_order_and_charts(self):
        result = self.analyze([
            row("A", stock=250, status="NEW RELEASE (RECOMMEND)"),
            row("B", stock=100, status="NORMAL"),
            row("C", stock=90, status="EOL"),
        ])
        excel_bytes = self.analyzer.export_excel(result)
        workbook = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertEqual(workbook.sheetnames, SHEET_NAMES)
        self.assertNotIn("库存结构概要", workbook.sheetnames)
        for sheet_name in SHEET_NAMES[:3]:
            headers = [cell.value for cell in workbook[sheet_name][1]]
            self.assertLess(headers.index("可用库存存销比"), headers.index("全供应链存销比"))
            self.assertEqual(headers, DETAIL_COLUMNS)
        status_sheet = workbook["按产品状态汇总"]
        self.assertEqual(status_sheet["A2"].value, "NEW RELEASE (RECOMMEND)")
        self.assertEqual(type(status_sheet._charts[0]).__name__, "PieChart")
        for sheet_name in ["新品推荐L1", "正常渐退L1", "停产L1"]:
            self.assertEqual(type(workbook[sheet_name]._charts[0]).__name__, "BarChart")
            self.assertEqual(len(workbook[sheet_name].tables), 0)
        with zipfile.ZipFile(io.BytesIO(excel_bytes)) as archive:
            chart_xml = "".join(
                archive.read(name).decode("utf-8")
                for name in archive.namelist()
                if name.startswith("xl/charts/chart")
            )
        self.assertIn("按产品状态概览", chart_xml)

    def test_excel_has_one_filterable_group_detail_with_numeric_model_counts(self):
        rows = [
            row("A", stock=100, status="NEW RELEASE (RECOMMEND)", group="G1", group_desc="Group 1"),
            row("B", stock=100, status="NEW RELEASE (RECOMMEND)", group="G2", group_desc="Group 2"),
        ]
        prices = pd.DataFrame({"物料代码": ["A", "B"], "单价": [2, 2]})
        result = self.analyze(rows, prices)
        workbook = openpyxl.load_workbook(io.BytesIO(self.analyzer.export_excel(result)))
        sheet = workbook["产品组分类明细"]
        self.assertIn("L1–L4 是不同口径", sheet["A1"].value)
        self.assertEqual([sheet["A2"].value, sheet["B2"].value], ["产品状态类别", "产品组描述"])
        self.assertEqual(sheet["E2"].value, "型号数量")
        self.assertEqual(sheet.auto_filter.ref, "A2:F4")
        self.assertEqual(sheet["E3"].value, 1)
        self.assertEqual(sheet["E3"].number_format, "#,##0")
        self.assertEqual(len(workbook["新品推荐L1"].tables), 0)
        self.assertIn("Group 1", [cell.value for row in sheet for cell in row])
        self.assertIn("Group 2", [cell.value for row in sheet for cell in row])

    def test_each_l_sheet_sorts_by_product_group_share_and_pivot_reconciles(self):
        rows = [
            row("A", stock=400, status="NEW RELEASE (RECOMMEND)", group="10", group_desc="Ten", category="Category A"),
            row("B", stock=200, status="NEW RELEASE (RECOMMEND)", group="2", group_desc="Two", category="Category A"),
            row("C", stock=300, status="NEW RELEASE (RECOMMEND)", group="1", group_desc="One", category="Category B"),
            row("D", stock=400, status="Active", group="3", group_desc="Three", category="Category B"),
        ]
        result = self.analyze(rows)
        detail = self.analyzer.summarize_excel_status_trade(result.detail, "NEW RELEASE (RECOMMEND)", "国贸产品分类码")
        self.assertEqual(detail["产品组"].tolist(), ["10", "1", "2"])
        self.assertEqual(detail.columns[:2].tolist(), ["产品组", "产品组描述"])
        pivot = self.analyzer.build_status_group_composition(detail, "国贸产品分类码")
        self.assertEqual(pivot["产品组"].tolist(), ["10", "1", "2"])
        self.assertAlmostEqual(
            pivot.iloc[:, 2:].to_numpy().sum(), detail[SUMMARY_AMOUNT_COLUMN].sum()
        )
        workbook = openpyxl.load_workbook(io.BytesIO(self.analyzer.export_excel(result)))
        sheet = workbook["新品推荐L1"]
        self.assertEqual([sheet[f"A{row_number}"].value for row_number in (2, 3, 4)], ["10", "1", "2"])
        self.assertEqual([sheet["A1"].value, sheet["B1"].value], ["产品组", "产品组描述"])
        self.assertEqual(sheet["D2"].number_format, "#,##0.00")
        self.assertEqual(sheet["E1"].value, "组内金额占比")
        self.assertEqual(sheet["E2"].number_format, "0.00%")
        self.assertEqual(sheet["F2"].number_format, "#,##0")
        charts = sheet._charts
        self.assertEqual(len(charts), 1)
        self.assertEqual(charts[0].grouping, "percentStacked")
        self.assertTrue(sheet.column_dimensions["AE"].hidden)
        self.assertEqual(sheet["AG2"].value, "Ten（44.4%）")
        self.assertEqual(charts[0].x_axis.scaling.orientation, "maxMin")
        self.assertIsNone(sheet["A8"].value)

    def test_composition_chart_shows_every_positive_category(self):
        categories = list("ABCDEF")
        amounts = [40, 25, 20, 10, 3, 2]
        result = self.analyze([
            row(code, stock=amount, status="NEW", category=code)
            for code, amount in zip(categories, amounts)
        ])
        summary = self.analyzer.summarize_excel_status_trade(
            result.detail, "NEW RELEASE (RECOMMEND)", "国贸产品分类码"
        )
        pivot = self.analyzer.build_status_group_composition(summary, "国贸产品分类码")
        self.assertEqual(pivot.columns[2:].tolist(), categories)
        self.assertAlmostEqual(pivot.iloc[0]["F"], 2 * 2 / 1_000_000)
        self.assertNotEqual(composition_chart_color("A", 0), "#FF0000")

        exact_eighty = pd.DataFrame({
            "产品组": ["1", "1", "1"],
            "产品组描述": ["Group 1"] * 3,
            "国贸产品分类码": ["A", "B", "C"],
            SUMMARY_AMOUNT_COLUMN: [50.0, 30.0, 20.0],
        })
        exact_pivot = self.analyzer.build_status_group_composition(exact_eighty, "国贸产品分类码")
        self.assertEqual(exact_pivot.columns[2:].tolist(), ["A", "B", "C"])

        exported = self.analyzer.export_excel(result)
        workbook = openpyxl.load_workbook(io.BytesIO(exported))
        chart = workbook["新品推荐L1"]._charts[0]
        self.assertEqual(len(chart.series), len(categories))
        self.assertEqual(chart.legend.position, "b")
        self.assertIsNotNone(chart.legend.layout.manualLayout)
        with zipfile.ZipFile(io.BytesIO(exported)) as archive:
            chart_xml = archive.read("xl/charts/chart2.xml")
        chart_root = ET.fromstring(chart_xml)
        chart_ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
        plot_layout = chart_root.find(".//c:plotArea/c:layout/c:manualLayout", chart_ns)
        legend_layout = chart_root.find(".//c:legend/c:layout/c:manualLayout", chart_ns)
        self.assertIsNotNone(plot_layout)
        self.assertIsNotNone(legend_layout)
        plot_bottom = sum(float(plot_layout.find(f"c:{field}", chart_ns).attrib["val"]) for field in ("y", "h"))
        legend_top = float(legend_layout.find("c:y", chart_ns).attrib["val"])
        self.assertLess(plot_bottom, legend_top)
        self.assertIsNone(chart.title)
        self.assertEqual(
            chart.series[-1].graphicalProperties.solidFill.srgbClr,
            composition_chart_color("F", 5).lstrip("#"),
        )

    def test_excel_shows_missing_trade_level_and_uses_legacy_l1(self):
        result = self.analyze([row("A", status="NEW RELEASE (RECOMMEND)")])
        workbook = openpyxl.load_workbook(io.BytesIO(self.analyzer.export_excel(result)))
        self.assertEqual(workbook["新品推荐L1"]["A1"].value, "产品组")
        self.assertEqual(workbook["新品推荐L1"]["C1"].value, "国贸产品分类码")
        self.assertEqual(workbook["新品推荐L2"]["A2"].value, "源文件未提供国贸产品分类码L2")

    def test_abc_sheet_sorts_each_grade_by_product_group_share(self):
        source = replenishment([
            row("A", stock=100, status="NORMAL", group="10", group_desc="Ten"),
            row("B", stock=300, status="NORMAL", group="2", group_desc="Two"),
        ])
        source[("产品基本信息", "泛欧ABC")] = ["A", "A"]
        result = self.analyzer.analyze(
            source, pd.DataFrame({"物料代码": ["A", "B"], "单价": [2, 2]})
        )
        summary = self.analyzer.summarize_excel_status_inventory(result.detail)
        self.assertEqual(summary["产品组"].tolist(), ["2", "10"])
        workbook = openpyxl.load_workbook(io.BytesIO(self.analyzer.export_excel(result)))
        sheet = workbook["正常渐退ABC"]
        self.assertEqual([sheet["A2"].value, sheet["A3"].value], ["2", "10"])
        self.assertEqual(sheet["J3"].value, "Two（75.0%）")
        self.assertEqual(sheet._charts[0].x_axis.scaling.orientation, "maxMin")


if __name__ == "__main__":
    unittest.main()
